from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from openpyxl import load_workbook

from src.settings import RAW_DIR, PROCESSED_DIR, SUPPORTED_EXTENSIONS


def ensure_dirs() -> None:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)


def extract_text_from_docx(file_path: Path) -> str:
    doc = Document(file_path)
    blocks = []
    # Walk body order so each rate table stays beside its product/tenor heading.
    for element in doc.element.body.iterchildren():
        if element.tag.endswith('}p'):
            text = Paragraph(element, doc).text.strip()
            if text:
                blocks.append(text)
        elif element.tag.endswith('}tbl'):
            table = Table(element, doc)
            rows = [" | ".join(cell.text.strip().replace("\n", "; ")
                               for cell in row.cells) for row in table.rows]
            if rows:
                blocks.append("\n".join(rows))
    return "\n\n".join(blocks)


def extract_text_from_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    text_chunks: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_xlsx(file_path: Path) -> str:
    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets_text: list[str] = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                clean_row = [str(cell) if cell is not None else "" for cell in row]
                if any(clean_row):
                    rows.append(" | ".join(clean_row))
            if rows:
                sheets_text.append("\n".join(rows))
        return "\n".join(sheets_text)
    finally:
        wb.close()


def extract_text_from_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")


def infer_bank_and_category(file_path: Path) -> tuple[str, str, str]:
    path_parts = [p for p in file_path.parts if p not in {"data", "raw", "processed"}]
    bank = "Unknown"
    category = "Other Services"
    lng = "bilingual"

    if len(path_parts) >= 2:
        bank = path_parts[-2]

    name = file_path.stem.lower()
    if any(token in name for token in ["loan", "loans", "credit"]):
        category = "Loans"
    elif any(token in name for token in ["card", "cards", "debit", "prepaid", "credit card"]):
        category = "Cards"
    elif any(token in name for token in ["account", "deposit", "saving", "current", "fixed"]):
        category = "Accounts"
    elif any(token in name for token in ["remittance", "payment", "wallet", "pay", "wire"]):
        category = "Remittance and Payments"
    elif any(token in name for token in ["merchant", "pos", "agent"]):
        category = "Merchant Services"
    elif any(token in name for token in ["trade", "finance", "lc", "supply"]):
        category = "Trade Finance"
    elif any(token in name for token in ["mbanking", "mobile", "digital", "online"]):
        category = "mBanking"
    elif any(token in name for token in ["insurance", "health", "life", "rider"]):
        category = "Insurance"

    if any(token in name for token in ["english", "eng"]):
        lng = "en"
    elif any(token in name for token in ["burmese", "myanmar", "mm"]):
        lng = "mm"

    return bank, category, lng


def process_file(file_path: Path) -> dict[str, Any]:
    suffix = file_path.suffix.lower()
    if suffix == ".docx":
        text = extract_text_from_docx(file_path)
    elif suffix == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif suffix in {".xlsx", ".xls"}:
        text = extract_text_from_xlsx(file_path)
    elif suffix in {".txt", ".md", ".csv"}:
        text = extract_text_from_txt(file_path)
    else:
        text = ""

    bank, category, language = infer_bank_and_category(file_path)
    output = {
        "source": file_path.name,
        "path": str(file_path),
        "bank": bank,
        "category": category,
        "language": language,
        "text": text,
    }
    return output


def load_processed_documents() -> list[dict[str, Any]]:
    ensure_dirs()
    documents: list[dict[str, Any]] = []

    for file_path in RAW_DIR.rglob("*"):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                document = process_file(file_path)
                if document["text"].strip():
                    documents.append(document)
            except Exception:
                continue

    return documents
